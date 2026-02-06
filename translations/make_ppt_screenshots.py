###############################################################################
# Copyright (c) 2018-2026 Lawrence Livermore National Security, LLC.
# Produced at the Lawrence Livermore National Laboratory
#
# Written by J. Brodsky, J. Chavez, S. Czyz, G. Kosinovsky, V. Mozin,
#            S. Sangiorgio.
#
# RASE-support@llnl.gov.
#
# LLNL-CODE-2014600, LLNL-CODE-829509
#
# All rights reserved.
#
# This file is part of RASE.
#
# Permission is hereby granted, free of charge, to any person obtaining a copy of
# this software and associated documentation files (the "Software"), to deal in
# the Software without restriction, including without limitation the rights to
# use, copy, modify, merge, publish, distribute, sublicense, and/or sell copies
# of the Software, and to permit persons to whom the Software is furnished to do
# so, subject to the following conditions:
#
# The above copyright notice and this permission notice shall be included in all
# copies or substantial portions of the Software.
#
# THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR
# IMPLIED,INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY, FITNESS
# FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE AUTHORS OR
# COPYRIGHT HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER LIABILITY, WHETHER
# IN AN ACTION OF CONTRACT, TORT OR OTHERWISE, ARISING FROM, OUT OF OR IN
# CONNECTION WITH THE SOFTWARE OR THE USE OR OTHER DEALINGS IN THE SOFTWARE.
###############################################################################
""" This script generates a PowerPoint presentation with side-to-side screenshots of the GUI in two languages.
This is useful for comparing the GUI translations.
"""

from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_PARAGRAPH_ALIGNMENT
import os

folder = '../tests/screenshots'
lang_set = ['de_DE', 'es_ES', 'it_IT']

for lang in lang_set:

    # get list of files in folder
    images = [f for f in os.listdir(os.path.join(folder,'en_US')) if f.endswith('.png')]

    if not images:
        print(f'No images found in {folder}/en_US/')
        print('Remember to run the TestGUIScreenshots test first')
        exit(1)

    # Create a new Presentation
    prs = Presentation('empty_presentation.pptx')

    # Define a blank slide layout (typically layout index 6 is blank)
    blank_slide_layout = prs.slide_layouts[6]

    # check images filenames are present also in the 'lang' folders
    for image in images:
        if not os.path.exists(os.path.join(folder, lang, image)):
            print(f'Image {image} not found in {folder}/{lang}/')
            images.remove(image)

    for image in images:
        # Add a new blank slide
        slide = prs.slides.add_slide(blank_slide_layout)

        # Add a textbox with the dialog name
        textbox = slide.shapes.add_textbox(left=Inches(0.1), top=Inches(0.05), width=Inches(7), height=Inches(0.4))
        textbox.text_frame.text = image.split('.')[0].split('_')[1]   # assume files are named as screenshot_<dialog_name>.png

        textbox2 = slide.shapes.add_textbox(left=Inches(0.1), top=Inches(0.05), width=Inches(13), height=Inches(0.4))
        p = textbox2.text_frame.paragraphs[0]
        p.text = lang
        p.alignment = PP_PARAGRAPH_ALIGNMENT.RIGHT

        # Add pictures next to each other. The coordinates (left, top) and size (width, height) can be adjusted
        slide.shapes.add_picture(os.path.join(folder, 'en_US', image), left=Inches(0.1), top=Inches(0.4), width=Inches(6.5))
        slide.shapes.add_picture(os.path.join(folder, lang, image), left=Inches(6.68), top=Inches(0.4), width=Inches(6.5))

    # Save the presentation to a file
    prs.save(f'RASE_screenshots_en_US_{lang}.pptx')